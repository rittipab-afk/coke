# Builds PI_Tuner_Template.xlsx: a PI DataLink export sheet for PID Loop Tuner.
# The DataLink "Sampled Data" formula itself is created by the user once on a plant PC
# (it needs PI DataLink and the plant PI server), so this file only prepares the layout.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

F = "Arial"
bold = Font(name=F, bold=True)
norm = Font(name=F)
blue = Font(name=F, color="0000FF")
title = Font(name=F, bold=True, size=14)
muted = Font(name=F, color="666666", italic=True)
yellow = PatternFill("solid", start_color="FFFF00")
head_fill = PatternFill("solid", start_color="DDEBF7")
thin = Side(style="thin", color="BBBBBB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ── Setup
s = wb.active
s.title = "Setup"
s["A1"] = "ตั้งค่า PI Tuner"; s["A1"].font = title
s["A2"] = "แก้เฉพาะช่องสีเหลือง ช่องอื่นคำนวณให้เอง"; s["A2"].font = muted
rows = [
    (4, "PI server", "GCMPPISVR", "ชื่อ PI Data Server (ส่วน \\\\GCMPPISVR\\ ในชื่อ tag ของ DataLink)"),
    (5, "Area prefix", "3-CTA.2M.", "ส่วนหน้าชื่อ tag ของพื้นที่นั้น (รวมจุดท้าย) ปกติไม่ต้องเปลี่ยนถ้าอยู่ area เดิม"),
    (6, "Loop tag", "3AC1102B", "ชื่อ loop เปลี่ยนช่องนี้ช่องเดียวเมื่อจะดู loop อื่น"),
    (7, "PV suffix", ".PV", "ถ้าชื่อ tag ใน PI ต่างจากนี้ แก้ได้"),
    (8, "SV suffix", ".SV", ""),
    (9, "MV suffix", ".MV", ""),
    (10, "MODE suffix", ".MODE", "ถ้าไม่มี tag MODE ใน PI ปล่อยไว้ได้ แอปยังใช้งานได้"),
    (12, "Start time", "*-2h", "PI time: *-2h = ย้อนหลัง 2 ชั่วโมง"),
    (13, "End time", "*", "* = ตอนนี้"),
    (14, "Interval", "1s", "flow/pressure 1s · temperature 5s · level 1–5s"),
]
for r, label, val, note in rows:
    s.cell(r, 1, label).font = bold
    c = s.cell(r, 2, val); c.font = blue; c.fill = yellow; c.border = box
    s.cell(r, 3, note).font = muted
s["A16"] = "ชื่อ tag เต็ม (คำนวณให้)"; s["A16"].font = bold
for i, (lab, ref) in enumerate([("PV", "B7"), ("SV", "B8"), ("MV", "B9"), ("MODE", "B10")]):
    r = 17 + i
    s.cell(r, 1, lab).font = norm
    # e.g. \\GCMPPISVR\3-CTA.2M.3AC1102B.PV  (Excel strings have no escapes: "\\" is two backslashes)
    c = s.cell(r, 2, '="\\\\"&$B$4&"\\"&$B$5&$B$6&' + ref); c.font = norm; c.border = box
s["A22"] = "สี:"; s["A22"].font = bold
s["B22"] = "ช่องที่แก้ได้"; s["B22"].fill = yellow; s["B22"].font = blue
s["C22"] = "ค่าในช่องเหลืองตอนนี้เป็นตัวอย่าง (loop 3AC1102B ที่ใช้ทดสอบ)"; s["C22"].font = muted
s.column_dimensions["A"].width = 24
s.column_dimensions["B"].width = 40
s.column_dimensions["C"].width = 70

# ── Data (PI DataLink output goes to A2)
d = wb.create_sheet("Data")
d["A1"] = "Timestamp"
for col, r in zip("BCDE", range(17, 21)):
    d[f"{col}1"] = f"=Setup!$B${r}"
for col in "ABCDE":
    d[f"{col}1"].font = bold; d[f"{col}1"].fill = head_fill; d[f"{col}1"].border = box
    d.column_dimensions[col].width = 34 if col != "A" else 22
d["A2"].comment = Comment("PI DataLink → Sampled Data ใส่ผลลัพธ์ที่เซลล์นี้ (A2) ดูขั้นตอนในแท็บ HowTo", "PI Tuner")
d["G1"] = "← แถว 1 คือชื่อ tag (ดึงจากแท็บ Setup) ข้อมูลจาก PI DataLink เริ่มที่ A2"; d["G1"].font = muted
d.freeze_panes = "A2"

# ── HowTo
h = wb.create_sheet("HowTo")
h["A1"] = "วิธีใช้"; h["A1"].font = title
steps = [
    ("ครั้งแรก (ทำครั้งเดียว บนเครื่องที่มี PI DataLink)", None),
    ("1", "แท็บ Setup: ตรวจ PI server (B4), area prefix (B5), ชื่อ loop (B6) และ suffix ให้ได้ชื่อ tag เต็มตรงกับที่ใช้ใน DataLink เช่น \\\\GCMPPISVR\\3-CTA.2M.3AC1102B.MV"),
    ("2", "แท็บ Data: คลิกเซลล์ A2"),
    ("3", "แท็บ PI DataLink ของ Excel → Sampled Data"),
    ("4", "Data item(s): ลากเลือก Data!B1:E1 (ชื่อ tag 4 ตัว)"),
    ("5", "Start time: คลิกเซลล์ Setup!B12 · End time: Setup!B13 · Time interval: Setup!B14 (หรือพิมพ์ *-2h, *, 1s ตรงๆ ก็ได้)"),
    ("6", "ติ๊ก Show timestamps · Output cell: Data!A2 แล้วกด OK"),
    ("7", "Save ไฟล์นี้ไว้ เช่น PI_Tuner.xlsx"),
    ("", ""),
    ("ครั้งต่อไป", None),
    ("1", "เปิดไฟล์ แล้วแก้ชื่อ loop ที่ Setup!B6 (ถ้าเปลี่ยน area ให้แก้ B5 ด้วย)"),
    ("2", "ถ้าข้อมูลไม่อัปเดตเอง: กด F9 หรือ PI DataLink → Recalculate/Refresh"),
    ("3", "ไปแท็บ Data → Ctrl+A → Ctrl+C"),
    ("4", "เปิด pid-tuner.html แล้วกด Ctrl+V ที่หน้าไหนก็ได้"),
    ("", ""),
    ("หมายเหตุ", None),
    ("•", "ถ้าเปลี่ยนช่วงเวลาให้ยาวขึ้นแล้วข้อมูลไม่ครบ ให้ลบผลเดิมใน Data แล้วทำขั้นตอนครั้งแรกข้อ 2–6 ใหม่"),
    ("•", "ถ้า DataLink ใส่คอลัมน์เวลาแยกให้ทุก tag ไม่ต้องลบ แอปอ่านได้"),
    ("•", "ถ้า tag .SV หรือ .MODE ไม่มีใน PI (ช่องขึ้น error) ให้ค้นชื่อจริงด้วย PI DataLink → Search แล้วแก้ suffix ใน Setup"),
    ("•", "ไฟล์นี้ไม่มีสูตร PI DataLink มาให้ ต้องสร้างด้วยหน้าต่าง Sampled Data ตามข้อ 3–6 เพราะสูตรต้องสร้างบนเครื่องที่ต่อ PI ได้"),
]
r = 3
for a, b in steps:
    if b is None:
        h.cell(r, 1, a).font = bold
    else:
        h.cell(r, 1, a).font = norm
        c = h.cell(r, 2, b); c.font = norm; c.alignment = wrap
    r += 1
h.column_dimensions["A"].width = 10
h.column_dimensions["B"].width = 100

wb.move_sheet("HowTo", offset=-2)  # HowTo first
wb.active = 0
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel computes the tag-name formulas on open
wb.save("PI_Tuner_Template.xlsx")
print("saved")
