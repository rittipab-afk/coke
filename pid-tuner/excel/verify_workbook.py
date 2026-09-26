# Verify PI_Loop_Health.xlsx formulas against the web app core (tested reference).
# Needs LibreOffice (recalc) and Node. Usage: python3 verify_workbook.py [workdir]
import json
import os
import subprocess
import sys
from datetime import datetime, timezone

from openpyxl import load_workbook

import make_workbook as mw

HERE = os.path.dirname(os.path.abspath(__file__))
RECALC = os.environ.get("RECALC", "/root/.claude/skills/synced/20a1a2ee-91b1-4af6-abb8-4fecc57d71f8_ce0049bc-51c9-4ae9-af9d-bd2d445f84eb/xlsx/scripts/recalc.py")
WORK = sys.argv[1] if len(sys.argv) > 1 else "/tmp/verify_wb"
os.makedirs(WORK, exist_ok=True)

# smaller sheets for speed (formulas are identical per row)
mw.MAX_ROWS = 4000
mw.R1 = mw.R0 + mw.MAX_ROWS - 1

fails = []


def check(cond, msg):
    print(("ok   " if cond else "FAIL ") + msg)
    if not cond:
        fails.append(msg)


def recalc(path):
    r = subprocess.run([sys.executable, RECALC, path, "300"], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    res = json.loads(r.stdout)
    if res.get("status") != "success":
        raise SystemExit(f"recalc failed for {path}: {r.stdout}")
    return res


def values(path, names):
    wb = load_workbook(path, data_only=True)
    out = {}
    for n in names:
        sheet, ref = names[n].split("!")
        out[n] = wb[sheet][ref.replace("$", "")].value
    return wb, out


ref = json.loads(subprocess.run(["node", os.path.join(HERE, "verify_ref.mjs")], capture_output=True, text=True, check=True).stdout)


def near(a, b, rel, absol=1e-9):
    return a is not None and b is not None and abs(a - b) <= max(abs(b) * rel, absol)


# ── loop health datasets
for ds in ref["datasets"]:
    rows = []
    for i in range(len(ds["t"])):
        ts = datetime.fromtimestamp(ds["t"][i] / 1000, tz=timezone.utc).replace(tzinfo=None)
        mode = ds["mode"][i] if ds["mode"] else None
        rows.append((ts, ds["pv"][i], ds["sp"][i], ds["op"][i], mode))
    path = os.path.join(WORK, ds["name"].replace(".csv", "") + ".xlsx")
    table = [("TEST", ds["loopType"], ds["sl"], ds["sh"], "u", None, None, None, "")]
    names = mw.build(path, data=rows, loop={"tag": "TEST", "table": table})
    res = recalc(path)
    check(res["total_errors"] == 0, f"{ds['name']}: no formula errors ({res['total_errors']})")
    wb, v = values(path, names)
    js = ds["js"]
    tag = ds["name"]
    for key, jkey, tol in (("ErrStd", "errStd", 0.01), ("ErrMean", "errMean", 0.01), ("OpLowPct", "opLowPct", 0.01), ("OpHighPct", "opHighPct", 0.01)):
        if js[jkey] is None:
            check(v[key] in (None, ""), f"{tag}: {key} empty like web app")
        else:
            check(near(v[key], js[jkey], tol, 1e-6), f"{tag}: {key} {v[key]} vs {js[jkey]:.6g}")
    if js["noise"] is not None:
        check(near(v["Noise"], js["noise"], 0.01, 1e-9), f"{tag}: Noise {v['Noise']} vs {js['noise']:.6g}")
    check(near(v["FlatFrac"], js["flatFrac"], 0.02, 0.005) and near(v["LinFrac"], js["linFrac"], 0.02, 0.005),
          f"{tag}: compression flat {v['FlatFrac']:.3f}/{js['flatFrac']:.3f} lin {v['LinFrac']:.3f}/{js['linFrac']:.3f}")
    check(bool(v["Compressed"]) == js["compressed"], f"{tag}: Compressed {v['Compressed']} vs {js['compressed']}")
    check(bool(v["SpMoving"]) == js["spMoving"], f"{tag}: SpMoving {v['SpMoving']} vs {js['spMoving']}")
    if js["errStd"] is not None:
        check(bool(v["Oscillating"]) == js["oscillating"], f"{tag}: Oscillating {v['Oscillating']} vs {js['oscillating']} (NZC {v['NZC']}, r {v['Regularity']}, peak {v['AcfPeak']})")
        if js["oscillating"] and v["Oscillating"]:
            check(near(v["MeanP"], js["period"], 0.15), f"{tag}: period {v['MeanP']:.1f} vs {js['period']:.1f} s")
    hs = wb["Health"]
    print(f"      Health: {hs['A6'].value} | {hs['A7'].value} | {hs['A9'].value}")
    check(all(hs[f"C{r}"].value not in (None, "") for r in range(18, 28)), f"{tag}: Health status column filled")

# ── tuning
for i, tc in enumerate(ref["tuning"]):
    path = os.path.join(WORK, f"tuning_{i}.xlsx")
    table = [("TEST", tc["loopType"], 0, 100, "u", 150, 10, 0, "")]
    tun = {"model": tc["model"], "Kp": tc.get("Kp"), "tau": tc.get("tau"), "theta": tc["theta"], "Ki": tc.get("Ki"), "Ts": tc["Ts"], "method": tc["method"], "lam": tc["lam"], "aggr": "ปกติ"}
    if tc["model"] == "Integrating":
        tun["Kp"] = None
        tun["tau"] = None
    names = mw.build(path, data=[], loop={"tag": "TEST", "table": table}, tuning=tun)
    res = recalc(path)
    check(res["total_errors"] == 0, f"tuning {i}: no formula errors")
    wb = load_workbook(path, data_only=True)
    tu = wb["Tuning"]
    PB, TI, TD = tu["B44"].value, tu["B45"].value, tu["B46"].value
    e = tc["expect"]
    ok = near(PB, e["PB"], 1e-6) and near(TI, e["TI"], 1e-6) and near(TD or 0, e["TD"] or 0, 1e-6, 1e-9)
    check(ok, f"tuning {i} {tc['model']} {tc['method']} ({tc['loopType']}): PB {PB} vs {e['PB']:.6g}, TI {TI} vs {e['TI']:.6g}, TD {TD} vs {e['TD']:.6g} [{tu['B48'].value}]")

print(f"\n{len(fails)} failure(s)")
sys.exit(1 if fails else 0)
